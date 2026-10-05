"""Word and PowerPoint extraction. A change anywhere a person would read must show up."""
import io

import pytest

docx = pytest.importorskip("docx")
pptx = pytest.importorskip("pptx")
from pptx.util import Inches  # noqa: E402

from canvas_compare import loader, parse_office  # noqa: E402
from canvas_compare.parse_office import extract_docx_text, extract_pptx_text  # noqa: E402
from conftest import manifest, modified_titles  # noqa: E402


def docx_bytes(build) -> bytes:
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_bytes(build) -> bytes:
    p = pptx.Presentation()
    build(p, p.slides.add_slide(p.slide_layouts[6]))
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def test_docx_table_text_is_extracted_in_order():
    def build(d):
        d.add_paragraph("Before")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Due", "Friday"
        d.add_paragraph("After")
    assert extract_docx_text(docx_bytes(build)).splitlines() == ["Before", "Due | Friday", "After"]


def test_docx_merged_cell_counts_once():
    def build(d):
        t = d.add_table(rows=1, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged"
        t.cell(0, 2).text = "End"
    assert extract_docx_text(docx_bytes(build)) == "Merged | End"


def test_docx_headers_and_footers():
    def build(d):
        d.add_paragraph("Body")
        d.sections[0].header.paragraphs[0].text = "Header text"
        d.sections[0].footer.paragraphs[0].text = "Footer text"
    assert extract_docx_text(docx_bytes(build)).splitlines() == ["Body", "[Header] Header text", "[Footer] Footer text"]


def test_pptx_table_group_and_notes():
    def build(p, slide):
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "In a group"
        slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table.cell(0, 1).text = "Cell"
        slide.notes_slide.notes_text_frame.text = "Mention the deadline"
    lines = extract_pptx_text(pptx_bytes(build)).splitlines()
    assert "In a group" in lines and "| Cell" in " ".join(lines) and "[Notes] Mention the deadline" in lines


def test_table_only_change_is_reported_as_modified(make_imscc, compare):
    """The original bug: a changed table cell compared as unchanged."""
    def build(name, cell):
        def doc(d):
            t = d.add_table(rows=1, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Due", cell
        return make_imscc(name, {"imsmanifest.xml": manifest(), "web_resources/h.docx": docx_bytes(doc)})
    report, *_ = compare(build("old", "Friday"), build("new", "Monday"))
    assert modified_titles(report) == ["[Document] web_resources/h.docx"]


def test_corrupt_office_file_warns_and_is_left_out(make_imscc, compare):
    """Two different corrupt files must not both read as the same 'error' text and so as unchanged."""
    def build(name, junk):
        return make_imscc(name, {"imsmanifest.xml": manifest(), "web_resources/bad.docx": junk})
    report, old, new = compare(build("old", b"garbage one"), build("new", b"other garbage"))
    assert old.text_items == {} and new.text_items == {}
    assert any("could not read web_resources/bad.docx" in w for w in old.warnings)


def test_missing_office_library_is_recorded_in_the_warnings(make_imscc, monkeypatch):
    """A saved report must say documents were skipped, not look like a complete comparison."""
    monkeypatch.setattr(loader, "HAVE_DOCX", False)
    monkeypatch.setattr(parse_office, "HAVE_DOCX", False)
    path = make_imscc("one", {"imsmanifest.xml": manifest(),
                              "web_resources/a.docx": docx_bytes(lambda d: d.add_paragraph("x"))})
    course = loader.load_course(path, label="old")
    assert course.skipped_office == 1 and course.text_items == {}
    assert any("1 .docx/.pptx file(s) were not compared because python-docx is not installed" in w
               for w in course.warnings)
