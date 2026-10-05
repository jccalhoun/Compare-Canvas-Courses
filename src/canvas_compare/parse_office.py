"""
Optional .docx/.pptx text extraction. python-docx and python-pptx are
soft dependencies — a course export can be compared without them, just
with office documents skipped (see loader.py, which checks HAVE_DOCX/
HAVE_PPTX before deciding whether to skip a file entirely).
"""

from __future__ import annotations

import io

try:
    import docx as _docx
    from docx.table import Table as _Table
    from docx.text.paragraph import Paragraph as _Paragraph
except ImportError:
    _docx = None

try:
    from pptx import Presentation as _Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE as _SHAPE_TYPE
except ImportError:
    _Presentation = None

# Public availability flags — this is the one place cli.py/loader.py
# should check for these, rather than reaching into the underscore-
# prefixed _docx/_Presentation names directly from another module.
HAVE_DOCX = _docx is not None
HAVE_PPTX = _Presentation is not None


def missing_optional_deps() -> list[str]:
    """Names (for pip install) of optional office packages not installed."""
    missing = []
    if not HAVE_DOCX:
        missing.append("python-docx")
    if not HAVE_PPTX:
        missing.append("python-pptx")
    return missing


def _cell_text(text: str) -> str:
    return " ".join(text.split())


def _docx_table_lines(table) -> list[str]:
    """One line per non-empty row, cells joined with " | " (a merged cell
    counts once), then any table nested inside a cell."""
    lines, nested = [], []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            # A merged cell is repeated once per column it spans. Track the cell
            # elements themselves rather than id()s, so correctness never
            # depends on the objects staying alive for the whole loop.
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            cells.append(_cell_text(cell.text))
            nested.extend(cell.tables)
        if any(cells):
            lines.append(" | ".join(cells))
    for inner in nested:
        lines.extend(_docx_table_lines(inner))
    return lines


def extract_docx_text(raw_bytes: bytes) -> str:
    """
    Text of a .docx, in document order: paragraphs, table rows, then headers
    and footers. Raises if the file can't be read; the loader turns that into
    a warning and leaves the file out. (Returning an error message as the
    "text" would make two different corrupt files compare as identical, and so
    as "unchanged".) Not read: text boxes and drawings.
    """
    if _docx is None:
        raise RuntimeError("python-docx is not installed")
    doc = _docx.Document(io.BytesIO(raw_bytes))
    lines = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = _Paragraph(child, doc).text
            if text.strip():
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            lines.extend(_docx_table_lines(_Table(child, doc)))
    extras = []
    for section in doc.sections:
        for label, part in (("Header", section.header), ("Footer", section.footer)):
            if part.is_linked_to_previous:
                continue
            text = " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            if text and f"[{label}] {text}" not in extras:
                extras.append(f"[{label}] {text}")
    return "\n".join(lines + extras)


def _pptx_shape_lines(shape) -> list[str]:
    if shape.shape_type == _SHAPE_TYPE.GROUP:
        return [line for inner in shape.shapes for line in _pptx_shape_lines(inner)]
    if getattr(shape, "has_table", False) and shape.has_table:
        lines = []
        for row in shape.table.rows:
            cells = [_cell_text(c.text_frame.text) for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        return lines
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        return [shape.text_frame.text.strip()]
    return []


def extract_pptx_text(raw_bytes: bytes) -> str:
    """
    All slide text: text boxes (including inside groups), table rows, and
    speaker notes. Raises if unreadable, for the same reason as .docx.
    """
    if _Presentation is None:
        raise RuntimeError("python-pptx is not installed")
    prs = _Presentation(io.BytesIO(raw_bytes))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            lines.extend(_pptx_shape_lines(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[Notes] {notes}")
    return "\n".join(lines)
